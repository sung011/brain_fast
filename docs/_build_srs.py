#!/usr/bin/env python3
"""예시 PDF와 같은 형식의 요구사항 정의서(.docx)를 생성한다."""

from __future__ import annotations

from pathlib import Path

import math

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager, patches
from matplotlib.lines import Line2D
from matplotlib.patches import FancyArrowPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = Path(__file__).resolve().parent
DIAGRAM_PATH = OUT_DIR / "usecase_diagram.png"
SCREEN_FLOW_PATH = OUT_DIR / "screen_flow.png"
DOCX_PATH = ROOT / "1.3._요구사항_정의서.docx"

FONT = "Apple SD Gothic Neo"
NAVY = "1B365D"
NAVY_RGB = RGBColor(0x1B, 0x36, 0x5D)
TEAL = "1A6B73"
HEADER_FILL = "1B365D"
ROW_ALT = "F4F7FA"
LABEL_FILL = "E8EEF4"


def font_prop(size: float, weight: str = "regular"):
    path = "/System/Library/Fonts/AppleSDGothicNeo.ttc"
    return font_manager.FontProperties(fname=path, size=size, weight=weight)


def draw_stick(ax, x: float, y: float, label: str, fp):
    """y는 몸통 중앙. 오른팔 끝을 연결점으로 쓴다."""
    head_r = 0.23
    head_y = y + 1.08
    ax.add_patch(patches.Circle((x, head_y), head_r, fill=False, lw=1.4, ec="black", zorder=5))
    ax.plot([x, x], [head_y - head_r, y + 0.12], color="black", lw=1.4, zorder=5)
    ax.plot([x - 0.40, x + 0.40], [y + 0.58, y + 0.58], color="black", lw=1.4, zorder=5)
    ax.plot([x, x - 0.34], [y + 0.12, y - 0.50], color="black", lw=1.4, zorder=5)
    ax.plot([x, x + 0.34], [y + 0.12, y - 0.50], color="black", lw=1.4, zorder=5)
    ax.text(x, y - 0.78, label, ha="center", va="top", fontproperties=fp, color="black")
    return (x + 0.40, y + 0.58)


def draw_uc(ax, xy, text, fp, w=2.45, h=0.78):
    cx, cy = xy
    ax.add_patch(
        patches.Ellipse(
            (cx + 0.08, cy - 0.08),
            w,
            h,
            facecolor="#bdbdbd",
            edgecolor="none",
            alpha=0.40,
            zorder=3,
        )
    )
    ax.add_patch(
        patches.Ellipse(xy, w, h, facecolor="white", edgecolor="black", lw=1.25, zorder=4)
    )
    ax.text(cx, cy, text, ha="center", va="center", fontproperties=fp, color="black", zorder=5)
    return (cx, cy, w / 2, h / 2)


def ellipse_edge(ell, toward):
    cx, cy, rx, ry = ell
    x0, y0 = toward
    dx, dy = x0 - cx, y0 - cy
    if dx == 0 and dy == 0:
        return cx + rx, cy
    t = 1.0 / math.sqrt((dx / rx) ** 2 + (dy / ry) ** 2)
    return cx + dx * t, cy + dy * t


def ellipse_at_angle(ell, deg):
    cx, cy, rx, ry = ell
    rad = math.radians(deg)
    return cx + rx * math.cos(rad), cy + ry * math.sin(rad)


def assoc(ax, p0, ell):
    x2, y2 = ellipse_edge(ell, p0)
    ax.add_line(Line2D([p0[0], x2], [p0[1], y2], color="black", lw=1.2, zorder=3))


def dashed_arrow(ax, src, dst, label, fp, label_off=(0.0, 0.22), src_angle=None, dst_angle=None):
    x1, y1 = ellipse_at_angle(src, src_angle) if src_angle is not None else ellipse_edge(src, dst[:2])
    x2, y2 = ellipse_at_angle(dst, dst_angle) if dst_angle is not None else ellipse_edge(dst, src[:2])
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=12,
            linestyle=(0, (4, 2.5)),
            linewidth=1.15,
            color="black",
            zorder=3,
        )
    )
    if label:
        mx = (x1 + x2) / 2 + label_off[0]
        my = (y1 + y2) / 2 + label_off[1]
        ax.text(
            mx,
            my,
            label,
            ha="center",
            va="center",
            fontproperties=fp,
            color="black",
            zorder=6,
            bbox={"boxstyle": "square,pad=0.08", "facecolor": "white", "edgecolor": "none"},
        )


def generalization(ax, child, parent, parent_angle):
    x2, y2 = ellipse_at_angle(parent, parent_angle)
    pcx, pcy = parent[0], parent[1]
    ox, oy = x2 - pcx, y2 - pcy
    length = math.hypot(ox, oy) or 1.0
    ux, uy = ox / length, oy / length
    size = 0.17
    bx, by = x2 + ux * size * 1.7, y2 + uy * size * 1.7
    px, py = -uy, ux
    x1, y1 = ellipse_edge(child, (bx, by))
    ax.add_line(Line2D([x1, bx], [y1, by], color="black", lw=1.15, zorder=3))
    ax.add_patch(
        patches.Polygon(
            [
                (x2, y2),
                (bx + px * size, by + py * size),
                (bx - px * size, by - py * size),
            ],
            closed=True,
            facecolor="white",
            edgecolor="black",
            lw=1.15,
            zorder=4,
        )
    )


def build_diagram() -> None:
    plt.rcParams["axes.unicode_minus"] = False
    w, h = 16.6, 12.8
    fig, ax = plt.subplots(figsize=(w, h), dpi=190)
    ax.set_xlim(0, w)
    ax.set_ylim(0, h)
    ax.set_aspect("equal")
    ax.axis("off")
    fig.patch.set_facecolor("#f4f6f8")
    ax.set_facecolor("#f4f6f8")

    for gx in np.arange(0, w + 0.01, 0.32):
        ax.axvline(gx, color="#d7dee6", lw=0.55, zorder=0)
    for gy in np.arange(0, h + 0.01, 0.32):
        ax.axhline(gy, color="#d7dee6", lw=0.55, zorder=0)

    actor_fp = font_prop(11)
    uc_fp = font_prop(8.2)
    sys_fp = font_prop(12)
    stereo_fp = font_manager.FontProperties(
        fname="/System/Library/Fonts/Helvetica.ttc", size=8.0
    )

    ax.add_patch(
        patches.Rectangle(
            (2.58, 0.28),
            13.70,
            12.20,
            facecolor="white",
            edgecolor="black",
            lw=1.55,
            zorder=1,
        )
    )
    ax.text(
        9.45,
        12.08,
        "메디스캔노트",
        ha="center",
        va="center",
        fontproperties=sys_fp,
        color="black",
        zorder=2,
    )

    guest_hand = draw_stick(ax, 1.18, 10.15, "비회원", actor_fp)
    learner_hand = draw_stick(ax, 1.18, 6.20, "학습자", actor_fp)
    admin_hand = draw_stick(ax, 1.18, 2.05, "관리자", actor_fp)

    signup = draw_uc(ax, (5.55, 11.20), "회원가입", uc_fp, w=2.30, h=0.74)
    login = draw_uc(ax, (5.55, 9.85), "로그인", uc_fp, w=2.20, h=0.74)
    glossary = draw_uc(ax, (5.55, 8.15), "의학용어 사전 조회", uc_fp, w=2.75, h=0.76)
    study = draw_uc(ax, (5.55, 6.50), "의료영상 학습", uc_fp, w=2.55, h=0.80)
    ai = draw_uc(ax, (5.55, 4.80), "AI 영상 분석", uc_fp, w=2.45, h=0.76)
    review = draw_uc(ax, (5.55, 3.50), "오답노트", uc_fp, w=2.30, h=0.74)
    stats = draw_uc(ax, (5.55, 2.35), "학습 통계·리포트", uc_fp, w=2.70, h=0.74)
    member = draw_uc(ax, (5.55, 1.35), "회원 관리", uc_fp, w=2.30, h=0.70)
    content = draw_uc(ax, (5.55, 0.55), "학습 콘텐츠 관리", uc_fp, w=2.70, h=0.68)

    find_pw = draw_uc(ax, (9.85, 9.85), "아이디/비밀번호 찾기", uc_fp, w=2.90, h=0.74)
    brain = draw_uc(ax, (13.35, 8.30), "뇌 영상 학습", uc_fp, w=2.30, h=0.70)
    chest = draw_uc(ax, (13.35, 7.40), "흉부 영상 학습", uc_fp, w=2.30, h=0.70)
    abdomen = draw_uc(ax, (13.35, 6.50), "복부 영상 학습", uc_fp, w=2.30, h=0.70)
    knee = draw_uc(ax, (13.35, 5.60), "무릎 영상 학습", uc_fp, w=2.30, h=0.70)
    share = draw_uc(ax, (9.55, 4.80), "화면 공유/\n영상 불러오기", uc_fp, w=2.50, h=0.86)
    roi = draw_uc(ax, (13.35, 3.90), "ROI 위치 평가", uc_fp, w=2.30, h=0.70)
    retry = draw_uc(ax, (13.35, 3.15), "다시 풀기", uc_fp, w=2.15, h=0.66)

    assoc(ax, guest_hand, signup)
    assoc(ax, learner_hand, login)
    assoc(ax, learner_hand, glossary)
    assoc(ax, learner_hand, study)
    assoc(ax, learner_hand, ai)
    assoc(ax, learner_hand, review)
    assoc(ax, learner_hand, stats)
    assoc(ax, admin_hand, member)
    assoc(ax, admin_hand, content)

    for child, ang in ((brain, 38), (chest, 14), (abdomen, -10), (knee, -36)):
        generalization(ax, child, study, ang)

    dashed_arrow(ax, find_pw, login, "«extend»", stereo_fp, label_off=(0.0, 0.26))
    dashed_arrow(ax, roi, study, "«extend»", stereo_fp, label_off=(-1.20, 0.08))
    dashed_arrow(ax, ai, share, "«include»", stereo_fp, label_off=(0.0, 0.28))
    dashed_arrow(ax, retry, review, "«extend»", stereo_fp, label_off=(0.0, 0.22))

    fig.savefig(DIAGRAM_PATH, dpi=190, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)


def set_run_font(run, size=11, bold=False, color=None, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_p(container, text, *, size=11, bold=False, align="left", space_after=8, space_before=0, color=None, first_line=0):
    p = container.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    if first_line:
        pf.first_line_indent = Cm(first_line)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def shade(cell, fill: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_borders(cell, color="1B365D", sz="8"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, **sides):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for key, val in sides.items():
        node = OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def v_align(cell, val="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), val)
    tcPr.append(v)


def write_cell(cell, text, *, size=10, bold=False, align="left", color=None, fill=None, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=RGBColor(255, 255, 255) if white else color)
    if fill:
        shade(cell, fill)
    set_borders(cell)
    set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
    v_align(cell, "center")


def add_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run("요구사항 정의서")
    set_run_font(run, size=10, bold=True, color=NAVY_RGB)
    hp.paragraph_format.space_after = Pt(2)

    # header bottom border
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), NAVY)
    pBdr.append(bottom)
    pPr.append(pBdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # page number field
    run1 = fp.add_run("- ")
    set_run_font(run1, size=9, color=NAVY_RGB)

    def add_fld(paragraph, instr):
        r1 = paragraph.add_run()
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        r1._r.append(fld1)
        r2 = paragraph.add_run()
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = instr
        r2._r.append(instr_el)
        r3 = paragraph.add_run()
        fld3 = OxmlElement("w:fldChar")
        fld3.set(qn("w:fldCharType"), "end")
        r3._r.append(fld3)
        for r in (r1, r2, r3):
            set_run_font(r, size=9, color=NAVY_RGB)

    add_fld(fp, " PAGE ")
    run2 = fp.add_run(" -")
    set_run_font(run2, size=9, color=NAVY_RGB)


def spec_table(doc, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    widths = (Cm(3.6), Cm(13.0))
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.width = widths[0]
        c1.width = widths[1]
        write_cell(c0, k, size=10, bold=True, align="center", fill=LABEL_FILL, color=NAVY_RGB)
        write_cell(c1, v, size=10, align="left")
        # keep event flow as multi-line
        if "\n" in v:
            c1.text = ""
            first = True
            for line_txt in v.split("\n"):
                if first:
                    p = c1.paragraphs[0]
                    first = False
                else:
                    p = c1.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.2
                run = p.add_run(line_txt)
                set_run_font(run, size=10)
            shade(c1, "FFFFFF")
            set_borders(c1)
            set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
            v_align(c1, "center")
    return table


def grid_table(doc, headers: list[str], data: list[list[str]], col_widths: list):
    table = doc.add_table(rows=1 + len(data), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = col_widths[i]
        write_cell(cell, h, size=10, bold=True, align="center", fill=HEADER_FILL, white=True)
    for r, row in enumerate(data):
        fill = "FFFFFF" if r % 2 == 0 else ROW_ALT
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.width = col_widths[c]
            align = "center" if c in (0, len(row) - 1) else "left"
            if c == 1 and len(headers) == 4:
                align = "center"
            write_cell(cell, val, size=9.5, align=align, fill=fill)
            if c == 2 and len(headers) >= 3:
                # description left
                write_cell(cell, val, size=9.5, align="left", fill=fill)
    return table


def heading(doc, text, size=16):
    add_p(doc, text, size=size, bold=True, color=NAVY_RGB, space_before=10, space_after=10)


def subheading(doc, text):
    add_p(doc, text, size=13, bold=True, color=NAVY_RGB, space_before=12, space_after=8)


def build_doc() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.7)
    add_header_footer(section)

    # ---- 표지 ----
    for _ in range(6):
        add_p(doc, "", size=12, space_after=0)
    add_p(doc, "요구사항 정의서", size=28, bold=True, align="center", color=NAVY_RGB, space_after=18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), TEAL)
    pBdr.append(bottom)
    pPr.append(pBdr)
    add_p(doc, "", size=10, space_after=16)
    add_p(doc, "주제명 : 메디스캔노트", size=14, bold=True, align="center", color=NAVY_RGB, space_after=4)
    add_p(doc, "AI 기반 의료영상 학습 및 판독 보조 플랫폼", size=13, align="center", color=RGBColor(0x1A, 0x6B, 0x73), space_after=28)
    add_p(doc, "2026. 09. 17.", size=13, align="center", color=NAVY_RGB, space_after=0)

    doc.add_page_break()

    # ---- I. 개요 ----
    heading(doc, "Ⅰ.  개  요")
    subheading(doc, "1.  문서의 목적")
    add_p(
        doc,
        "본 문서는 “메디스캔노트(AI 기반 의료영상 학습 및 판독 보조 플랫폼)” 개발의 요구사항을 정의한 문서이다.",
        align="justify",
        first_line=0.4,
    )
    add_p(
        doc,
        "이 요구사항 정의서는 의료 영상(뇌·흉부·복부·무릎, X-ray/CT/MRI) 분석, 화면 공유 ROI 분석, 학습 문제 풀이, 병변 위치 채점, 오답 노트, 의학용어 사전, 관리자 회원·학습 콘텐츠 관리를 포함한 서비스의 설계 문서를 작성하는 데 기초가 되며, 사용자 유스케이스를 기반으로 SW가 제공해야 할 기능 및 화면에서 필수적으로 구현되어야 할 요구사항에 대해 기술하고 정의한다.",
        align="justify",
        first_line=0.4,
    )
    add_p(
        doc,
        "본 문서는 가능한 구체적이며 간결하게 표현되어야 하고 추후 시험이 가능해야 하며, 본 문서를 사용하는 대상은 본 과제를 기획하는 기획자, SW를 개발하는 개발자 등이며, 본 과제의 요구사항 도출 및 개발 과정에서 본 문서를 활용할 수 있도록 한다.",
        align="justify",
        first_line=0.4,
    )

    subheading(doc, "2.  요구사항 및 문서의 범위")
    add_p(
        doc,
        "본 문서에서는 유스케이스 및 기능·비기능 요구사항의 기술을 그 범위로 한다. 대상 시스템은 FastAPI 웹 애플리케이션, PostgreSQL(mediscan_note), Synology NAS 파일 저장, 부위별 의료영상 분류 모델, MedGemma 판독문 생성으로 구성된다. 지원 부위는 뇌·흉부·복부·무릎이며, 영상 종류는 X-ray·CT·MRI이다.",
        align="justify",
        first_line=0.4,
    )
    add_p(doc, "관련 액터는 다음과 같다.", align="justify", first_line=0.4)
    add_p(doc, "·  비회원 : 회원가입·로그인 전 사용자", size=10.5, space_after=4)
    add_p(doc, "·  학습자 : 메인 화면에서 의학용어 사전, 의료영상 학습, AI 영상 분석, 오답노트, 학습 통계를 사용한다. 로그인이 되어 있지 않으면 기능을 사용할 수 없고 로그인 화면으로 이동한다.", size=10.5, space_after=4)
    add_p(doc, "·  관리자 : 회원 및 학습 콘텐츠를 등록·수정·삭제하고 NAS 저장을 관리하는 사용자(mb_level=10)", size=10.5, space_after=4)
    add_p(doc, "·  시스템 : 분석 엔진, 데이터베이스, NAS File Station, 오류 로그 모듈", size=10.5, space_after=12)

    # ---- II. 유스케이스 ----
    heading(doc, "Ⅱ.  유스케이스")
    subheading(doc, "1.  유스케이스 다이어그램")
    add_p(
        doc,
        "메디스캔노트의 비회원·학습자·관리자 유스케이스와 일반화(뇌·흉부·복부·무릎), «include»/«extend» 관계를 나타낸다.",
        size=10.5,
        align="justify",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(DIAGRAM_PATH), width=Cm(16.5))

    doc.add_page_break()
    subheading(doc, "2.  화면 구성도")
    add_p(
        doc,
        "실제 서비스에 필요한 화면과 기능을 하나의 사용자 흐름으로 정리한 구성도이다. 회원/계정, 메인, 학습 진입, 핵심 학습, AI 분석, 오답노트, 학습 관리 영역으로 구성된다.",
        size=10.5,
        align="justify",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(SCREEN_FLOW_PATH), width=Cm(16.5))

    doc.add_page_break()
    subheading(doc, "3.  유스케이스 명세")

    specs = [
        [
            ("유스케이스 이름", "로그인"),
            ("유스케이스 ID", "U_C_201"),
            ("관련 요구사항", "F_R_201, F_R_206, F_R_239"),
            ("우선순위", "상"),
            ("선행조건", "학습자가 메인 화면에 접속한 상태이다. 로그인하려면 회원 계정이 등록되어 있고 상태가 정상(N)이어야 한다."),
            ("관련 액터", "학습자"),
            (
                "이벤트 흐름",
                "1. 학습자가 서비스에 접속하면 첫 화면으로 메인 화면이 표시된다.\n"
                "2. 로그인되어 있지 않으면 메인 화면에서 어떤 기능도 수행할 수 없다.\n"
                "3. 메인 화면에서 기능을 선택하면 로그인 화면으로 이동한다.\n"
                "4. 학습자는 아이디(이메일)와 비밀번호를 입력한다.\n"
                "5. 시스템은 입력값을 검증하고 user_member 테이블에서 해당 계정을 조회한다.\n"
                "6. 시스템은 bcrypt로 저장된 비밀번호와 입력값을 비교한다.\n"
                "7. 인증에 성공하면 세션을 생성하고, 원래 선택했던 기능 또는 메인 화면으로 돌아간다.\n"
                "8. 인증에 실패하면 “아이디 또는 비밀번호가 올바르지 않습니다.”를 출력하며, 기능은 수행되지 않는다.",
            ),
            ("종료조건", "로그인 성공 후 요청 기능/메인 복귀, 또는 로그인 실패로 기능 차단"),
        ],
        [
            ("유스케이스 이름", "화면 공유 ROI 분석"),
            ("유스케이스 ID", "U_C_202"),
            ("관련 요구사항", "F_R_211, F_R_212, F_R_213, F_R_214, F_R_215, F_R_216, F_R_217, F_R_219, F_R_220"),
            ("우선순위", "상"),
            ("선행조건", "메인 화면에 접속할 수 있어야 한다. 로그인이 되어 있지 않으면 기능을 사용할 수 없고 로그인 화면으로 이동한다. 분류 모델은 최초 요청 시 로드될 수 있다."),
            ("관련 액터", "학습자"),
            (
                "이벤트 흐름",
                "1. 학습자는 메인 화면에서 AI 영상 분석(화면 공유 ROI 분석)을 선택한다.\n"
                "2. 로그인되어 있지 않으면 기능은 수행되지 않고 로그인 화면으로 이동한다.\n"
                "3. 로그인 후 부위(뇌/흉부/복부/무릎)와 영상 종류(X-ray/CT/MRI)를 선택한다.\n"
                "4. 사용자는 화면 공유를 시작하거나 영상을 불러오고, 분석할 영역을 ROI(네모 박스)로 지정한다.\n"
                "5. 시스템은 ROI 이미지를 잘라 /api/v1/analyze/screen-roi 로 전송한다.\n"
                "6. 이미지가 5MB를 초과하거나 디코딩에 실패하면 오류를 반환한다.\n"
                "7. 분석 엔진은 선택한 부위·영상 종류에 맞는 분류 모델로 이상 소견 확률을 산출한다.\n"
                "8. 시스템은 이상 부위 오버레이와 한글 병명·신뢰도를 반환한다.\n"
                "9. 결과는 “AI 보조 참고용이며 실제 의료 진단이 아니다”는 면책 문구와 함께 표시한다.\n"
                "10. NAS_UPLOAD_ON_ANALYZE가 켜져 있으면 원본·오버레이·결과 JSON을 NAS에 저장한다.",
            ),
            ("종료조건", "분류 결과·오버레이·면책 문구 표시"),
        ],
        [
            ("유스케이스 이름", "AI 판독문 생성"),
            ("유스케이스 ID", "U_C_203"),
            ("관련 요구사항", "F_R_214, F_R_217, F_R_218, F_R_239"),
            ("우선순위", "상"),
            ("선행조건", "메인 화면에서 분석을 진행할 수 있는 상태여야 한다. 로그인이 되어 있지 않으면 로그인 화면으로 이동한다. 화면 공유 ROI 분석이 완료되어 클래스 확률(probs)이 있어야 한다."),
            ("관련 액터", "학습자"),
            (
                "이벤트 흐름",
                "1. 학습자는 메인 화면을 거쳐 분석 결과 화면에서 판독문(설명) 생성을 요청한다.\n"
                "2. 로그인되어 있지 않으면 기능은 수행되지 않고 로그인 화면으로 이동한다.\n"
                "3. 로그인 후 시스템은 원본 이미지, 확률 배열, 오버레이를 /api/v1/analyze/screen-roi/explain 로 전송한다.\n"
                "4. MedGemma(기본: Hugging Face Inference API)가 원본·오버레이를 바탕으로 설명을 생성한다.\n"
                "5. API 호출이 실패하면 설정에 따라 템플릿 요약으로 대체할 수 있다.\n"
                "6. 시스템은 생성된 판독문과 면책 문구를 반환한다.",
            ),
            ("종료조건", "판독문 텍스트 반환"),
        ],
        [
            ("유스케이스 이름", "학습 문제 풀이"),
            ("유스케이스 ID", "U_C_204"),
            ("관련 요구사항", "F_R_221, F_R_222, F_R_223, F_R_224, F_R_225, F_R_239"),
            ("우선순위", "상"),
            ("선행조건", "메인 화면에 접속할 수 있어야 하며, study 테이블에 삭제되지 않은 학습 콘텐츠가 1건 이상 있어야 한다. 로그인이 되어 있지 않으면 기능을 사용할 수 없고 로그인 화면으로 이동한다."),
            ("관련 액터", "학습자"),
            (
                "이벤트 흐름",
                "1. 학습자는 메인 화면에서 의료영상 학습을 선택한다.\n"
                "2. 로그인되어 있지 않으면 기능은 수행되지 않고 로그인 화면으로 이동한다.\n"
                "3. 로그인 후 부위(뇌/흉부/복부/무릎)와 영상 종류(X-ray/CT/MRI)를 선택한다.\n"
                "4. 시스템은 study에서 조건에 맞는 문제를 무작위로 1건 조회한다.\n"
                "5. 화면에는 이미지·부위·모달리티만 보여주고 정답 병명(st_disease)은 숨긴다.\n"
                "6. 사용자는 병명을 입력하여 제출한다.\n"
                "7. 시스템은 제출값과 정답을 비교하여 정오를 반환한다.\n"
                "8. 제출 이후에만 이미 푼 문항을 제외한 다음 무작위 문제를 제공한다.\n"
                "9. 조건에 맞는 문제가 없으면 “학습 문제가 없습니다.”를 출력한다.",
            ),
            ("종료조건", "채점 결과 및 다음 문제 반환"),
        ],
        [
            ("유스케이스 이름", "ROI 위치 채점 및 오답 노트 저장"),
            ("유스케이스 ID", "U_C_205"),
            ("관련 요구사항", "F_R_213, F_R_214, F_R_215, F_R_226, F_R_227, F_R_228, F_R_234"),
            ("우선순위", "상"),
            ("선행조건", "메인 화면에서 학습을 진행할 수 있는 상태여야 한다. 로그인이 되어 있지 않으면 로그인 화면으로 이동한다. 학습 문제 이미지가 표시되어 있고 NAS 계정 설정이 완료되어 있어야 한다."),
            ("관련 액터", "학습자"),
            (
                "이벤트 흐름",
                "1. 학습자는 메인 화면에서 학습을 선택한 뒤, 영상에서 이상이 있다고 판단한 위치를 박스 또는 원(ROI)으로 그린다.\n"
                "2. 로그인되어 있지 않으면 기능은 수행되지 않고 로그인 화면으로 이동한다.\n"
                "3. 로그인 후 시스템은 이미지와 ROI 좌표를 /learning/roi-grade 로 전송한다.\n"
                "4. 분석 엔진이 이상 부위를 찾고, 사용자 ROI와의 IoU·재현율·중심 거리로 채점한다.\n"
                "5. IoU 0.40 또는 재현율 0.50 이상이면 정답(C), IoU 0.15·재현율 0.25·근접이면 부분정답(H),\n"
                "   그 외 또는 이상 부위가 없으면 오답(W)으로 판정한다.\n"
                "6. 이상이 있으면 색칠된 오버레이를 함께 반환한다.\n"
                "7. 사용자 ROI가 그려진 이미지를 NAS review 폴더에 업로드한다.\n"
                "8. review_node에 회원, 학습 문항, 병명, 이미지 경로, 채점 결과(C/H/W)를 저장한다.",
            ),
            ("종료조건", "채점 결과 반환 및 오답 노트 저장"),
        ],
        [
            ("유스케이스 이름", "의학용어 사전 조회"),
            ("유스케이스 ID", "U_C_206"),
            ("관련 요구사항", "F_R_229"),
            ("우선순위", "중"),
            ("선행조건", "메인 화면에 접속할 수 있어야 한다. 로그인이 되어 있지 않으면 기능을 사용할 수 없고 로그인 화면으로 이동한다."),
            ("관련 액터", "학습자"),
            (
                "이벤트 흐름",
                "1. 학습자는 메인 화면에서 의학용어 사전을 선택한다.\n"
                "2. 로그인되어 있지 않으면 기능은 수행되지 않고 로그인 화면으로 이동한다.\n"
                "3. 로그인 후 부위(region: brain, chest, abdomen, knee 등)를 지정한다.\n"
                "4. (선택) 영상 종류(type: xray, ct, mri)를 추가로 지정한다.\n"
                "5. 시스템은 medical_terms에서 해당 부위·모달리티(및 common) 용어를 조회한다.\n"
                "6. 한글명, 영문명, 정의, 태그, 카테고리를 목록으로 반환한다.\n"
                "7. region이 없거나 허용되지 않은 값이면 오류 메시지를 출력한다.",
            ),
            ("종료조건", "용어 목록 반환"),
        ],
        [
            ("유스케이스 이름", "회원 관리"),
            ("유스케이스 ID", "U_C_207"),
            ("관련 요구사항", "F_R_204, F_R_205, F_R_206, F_R_207, F_R_208, F_R_209, F_R_210"),
            ("우선순위", "상"),
            ("선행조건", "N/A"),
            ("관련 액터", "관리자"),
            (
                "이벤트 흐름",
                "1. 관리자는 /admin/user 에서 회원 목록을 조회한다.\n"
                "2. 회원가입 화면에서 아이디 중복을 확인한 뒤 이름, 생년월일, 학교, 전공, 비밀번호를 등록한다.\n"
                "3. 비밀번호는 평문으로 저장하지 않고 해시하여 저장한다.\n"
                "4. 관리자는 회원 이름·소속·권한(일반 1 / 최고관리자 10)을 수정할 수 있다.\n"
                "5. 관리자는 회원을 논리 삭제할 수 있다.\n"
                "6. 등록·수정·삭제 시 SSE로 열려 있는 목록 화면을 새로고침 없이 갱신한다.",
            ),
            ("종료조건", "회원 정보의 등록·수정·삭제 반영"),
        ],
        [
            ("유스케이스 이름", "학습 콘텐츠 관리"),
            ("유스케이스 ID", "U_C_208"),
            ("관련 요구사항", "F_R_230, F_R_231, F_R_232, F_R_233, F_R_234, F_R_235, F_R_236, F_R_237"),
            ("우선순위", "상"),
            ("선행조건", "NAS 연결 설정이 존재해야 한다."),
            ("관련 액터", "관리자"),
            (
                "이벤트 흐름",
                "1. 관리자는 /admin/study 에서 학습 목록을 조회한다.\n"
                "2. 학습 등록 화면에서 부위, 영상 종류, 병명, 이미지를 입력한다.\n"
                "3. 시스템은 부위·모달리티에 맞는 NAS 폴더(예: /stylesheets/assets/brain/CT)에 이미지를 업로드한다.\n"
                "4. study 테이블에 웹 경로와 메타데이터(st_part, st_modal, st_disease)를 저장한다.\n"
                "5. 관리자는 상세 화면에서 메타를 수정하거나 이미지를 교체할 수 있다.\n"
                "6. 관리자는 학습 데이터를 논리 삭제할 수 있다.\n"
                "7. 변경 사항은 SSE로 목록 화면에 실시간 반영된다.",
            ),
            ("종료조건", "학습 콘텐츠의 등록·수정·삭제 반영"),
        ],
        [
            ("유스케이스 이름", "로그아웃"),
            ("유스케이스 ID", "U_C_209"),
            ("관련 요구사항", "F_R_202"),
            ("우선순위", "하"),
            ("선행조건", "로그인한 상태여야 한다."),
            ("관련 액터", "학습자"),
            (
                "이벤트 흐름",
                "1. 학습자는 메인 화면에서 로그아웃을 선택한다.\n"
                "2. 시스템은 세션을 비운다.\n"
                "3. 메인 화면으로 돌아간다.\n"
                "4. 이후 메인 화면에서 기능을 선택하면 다시 로그인 화면으로 이동한다.",
            ),
            ("종료조건", "세션 종료 후 메인 화면 복귀"),
        ],
    ]

    for i, spec in enumerate(specs):
        spec_table(doc, spec)
        if i < len(specs) - 1:
            add_p(doc, "", size=6, space_after=10)

    doc.add_page_break()

    # ---- III. 기능 요구사항 ----
    heading(doc, "Ⅲ.  기능 요구사항")
    add_p(
        doc,
        "우선순위는 상(필수)·중(중요)·하(선택)로 구분한다. 상은 핵심 기능과, 그 기능을 쓰려면 반드시 필요한 선행 조건이다. 중은 품질·편의 기능이다. 하(선택)는 없어도 핵심 시나리오가 성립하는 기능이다. 시험 시 본 표의 설명으로 합격 여부를 판정한다.",
        size=10.5,
        align="justify",
        space_after=10,
    )

    func_rows = [
        ["F_R_201", "로그인", "서비스 첫 화면은 메인 화면이다. 로그인이 되어 있지 않으면 메인에서 어떤 기능을 선택해도 동작하지 않고 로그인 화면으로 이동해야 한다. 아이디·비밀번호로 로그인할 수 있어야 하며, 실패 시 원인을 추측할 수 없는 오류 메시지를 출력해야 한다.", "상"],
        ["F_R_202", "로그아웃", "로그인한 상태에서 로그아웃을 할 수 있어야 하며, 세션이 즉시 만료되고 메인 화면으로 돌아가야 한다. 이후 기능을 선택하면 다시 로그인 화면으로 이동해야 한다.", "하"],
        ["F_R_203", "관리자 기능 제공", "관리자는 회원 관리와 학습 콘텐츠 관리 기능을 사용할 수 있어야 한다.", "상"],
        ["F_R_204", "아이디 중복 확인", "회원가입 시 아이디(이메일) 사용 가능 여부를 조회할 수 있어야 한다. 이미 사용 중이면 가입을 막아야 한다.", "상"],
        ["F_R_205", "회원 등록", "아이디, 비밀번호, 이름(최대 5자 권장/최대 10자), 생년월일, 학교, 전공을 입력하여 회원을 등록할 수 있어야 한다.", "상"],
        ["F_R_206", "비밀번호 보호", "비밀번호는 평문으로 저장하지 않고 해시(bcrypt)하여 저장해야 한다. 조회 API 응답에 비밀번호를 포함하지 않아야 한다.", "상"],
        ["F_R_207", "회원 목록 조회", "관리자는 등록된 회원 목록(아이디, 이름, 생년월일, 학교, 전공, 권한, 상태)을 조회할 수 있어야 한다.", "상"],
        ["F_R_208", "회원 정보 수정", "관리자는 이름, 생년월일, 학교, 전공, 권한(1:일반 / 10:최고관리자)을 수정할 수 있어야 하며, 비밀번호는 입력한 경우에만 변경해야 한다.", "상"],
        ["F_R_209", "회원 삭제", "관리자는 회원을 삭제할 수 있어야 한다. 실제 물리 삭제 대신 논리 삭제(del_yn 등)를 사용해야 한다.", "상"],
        ["F_R_210", "회원 변경 실시간 반영", "회원 등록·수정·삭제 시 열려 있는 목록 화면이 새로고침 없이 갱신되어야 한다(SSE).", "중"],
        ["F_R_211", "화면 공유", "사용자는 브라우저 화면 공유로 의료영상(뇌·흉부·복부·무릎, X-ray/CT/MRI)을 캡처하여 분석 화면에 표시할 수 있어야 한다.", "상"],
        ["F_R_212", "ROI 영역 지정", "공유된 화면(스크롤된 영역 포함)에서 네모 박스로 분석 영역을 지정할 수 있어야 한다.", "상"],
        ["F_R_213", "이미지 크기 제한", "업로드·분석 이미지는 최대 5MB여야 하며, 초과 시 413 또는 안내 메시지를 출력해야 한다. 손상·빈 이미지는 거부해야 한다.", "상"],
        ["F_R_214", "의료영상 이상 소견 분류", "뇌·흉부·복부·무릎 의료영상(X-ray/CT/MRI) ROI에 대해 해당 부위의 이상 소견 확률을 산출하고 상위 소견을 제시해야 한다.", "상"],
        ["F_R_215", "이상 부위 오버레이", "요청 시 클래스별 색상으로 이상 부위를 표시한 PNG 오버레이와 위치 요약을 제공해야 한다.", "상"],
        ["F_R_216", "한글 병명 표시", "분류 결과는 영문 클래스명과 함께 해당 부위의 한글 병명(소견명)을 제공해야 한다.", "상"],
        ["F_R_217", "의료 면책 문구", "분석·판독 결과에는 “AI 보조 참고용이며 실제 의료 진단이 아니다. 최종 판단은 의사에게 있다.”는 문구가 포함되어야 한다.", "상"],
        ["F_R_218", "판독문 생성", "분류 확률과 원본·오버레이 이미지를 바탕으로 자연어 판독문(설명)을 생성할 수 있어야 한다. API 실패 시 템플릿 요약으로 대체할 수 있다.", "상"],
        ["F_R_219", "분석 결과 NAS 저장", "설정이 활성화된 경우 분석 원본, 오버레이, 결과 JSON을 NAS에 저장할 수 있어야 한다. 설정이 비어 있으면 분석을 막지 않고 저장 실패 정보만 반환한다.", "중"],
        ["F_R_220", "분석 모듈 헬스체크", "/api/v1/health 에서 분류 모델·판독문 모델 로드 여부, 오류, 사용 디바이스를 조회할 수 있어야 한다. 모델 미로드 시에도 관리자 페이지는 동작해야 한다.", "상"],
        ["F_R_221", "랜덤 학습 문제 조회", "study 테이블에서 학습 문제 1건을 무작위로 조회할 수 있어야 한다.", "상"],
        ["F_R_222", "부위·모달리티 필터", "부위(뇌/흉부/복부/무릎)와 영상 종류(X-ray/CT/MRI)로 문제를 필터할 수 있어야 하며, 숫자 코드와 영문·한글 별칭을 모두 허용해야 한다.", "상"],
        ["F_R_223", "정답 병명 숨김", "문제 조회 응답에는 정답 병명(st_disease)을 포함하지 않아야 한다. 정답은 제출 후에만 공개해야 한다.", "상"],
        ["F_R_224", "병명 제출·채점", "사용자가 입력한 병명과 정답을 비교하여 정오 여부를 반환해야 한다.", "상"],
        ["F_R_225", "다음 문제 제공", "제출 성공 후에만 다음 무작위 문제를 제공해야 한다. 이미 푼 문항(exclude_idxs)은 제외해야 한다.", "상"],
        ["F_R_226", "ROI 박스/원 제출", "사용자는 박스(x,y,width,height) 또는 원(cx,cy,radius)으로 병변 위치를 제출할 수 있어야 한다. 좌표는 정규화(0~1) 또는 픽셀을 지원해야 한다.", "상"],
        ["F_R_227", "위치 채점 기준", "사용자 ROI와 AI 이상 부위의 IoU·재현율·중심 거리로 정답/부분정답/오답을 판정해야 한다. 판정 수치(iou, recall 등)를 함께 반환해야 한다.", "상"],
        ["F_R_228", "오답 노트 저장", "ROI 채점 제출 시 사용자 ROI가 표시된 이미지를 NAS review 경로에 저장하고, review_node에 회원·문항·병명·채점코드(C/H/W)를 기록해야 한다.", "상"],
        ["F_R_229", "의학용어 사전 조회", "부위·영상 종류로 의학용어(한글명, 영문명, 정의, 태그)를 조회할 수 있어야 한다. 모달리티 지정 시 common 용어도 함께 반환해야 한다.", "중"],
        ["F_R_230", "학습 목록 조회", "관리자는 학습 콘텐츠 목록(부위, 모달리티, 병명, 이미지 경로)을 조회할 수 있어야 한다.", "상"],
        ["F_R_231", "학습 등록", "관리자는 부위, 영상 종류, 병명, 이미지를 등록할 수 있어야 한다. 이미지는 NAS에 업로드된 뒤 DB에는 웹 상대 경로만 저장해야 한다.", "상"],
        ["F_R_232", "학습 수정", "관리자는 학습 메타데이터를 수정할 수 있어야 하며, 이미지 파일이 있으면 NAS에 새로 올린 뒤 경로를 갱신해야 한다.", "상"],
        ["F_R_233", "학습 삭제", "관리자는 학습 데이터를 논리 삭제할 수 있어야 한다.", "상"],
        ["F_R_234", "NAS 이미지 업로드", "학습·오답 노트 이미지는 Synology File Station API로 업로드할 수 있어야 한다. NAS 미설정 시 503과 안내 메시지를 출력해야 한다.", "상"],
        ["F_R_235", "부위별 저장 경로", "이미지는 부위·모달리티 폴더 규칙으로 저장되어야 한다. 예) 뇌+CT, 흉부+X-ray, 복부+CT, 무릎+MRI", "상"],
        ["F_R_236", "학습 변경 실시간 반영", "학습 등록·수정·삭제 시 열려 있는 목록 화면이 새로고침 없이 갱신되어야 한다(SSE).", "중"],
        ["F_R_237", "NAS 연결 확인", "관리자는 NAS 로그인 가능 여부를 조회할 수 있어야 한다.", "상"],
        ["F_R_238", "DB 헬스체크", "/health/db 로 PostgreSQL 연결 여부와 데이터베이스 이름을 확인할 수 있어야 한다.", "상"],
        ["F_R_239", "에러 메시지 출력", "입력값 오류, 인증 실패, 네트워크·NAS·모델 오류 등 다양한 이유로 오류가 발생하는 경우, 해당 에러 메시지를 출력해야 한다.", "상"],
        ["F_R_240", "오류 로그 기록", "입력값 검증 실패(422), HTTP 예외, 서버 예외(500)는 logs/YYYY-MM-DD/error.log 에 요청 정보와 함께 기록되어야 한다.", "상"],
        ["F_R_241", "로그 보관 주기", "한국 시간 기준으로 7일이 지난 날짜 로그 폴더는 자동 삭제되어야 한다. 서버 기동 시와 1시간마다 검사해야 한다.", "중"],
        ["F_R_242", "컨테이너 실행", "애플리케이션과 PostgreSQL은 Docker Compose로 함께 기동할 수 있어야 한다. 앱은 8000, DB는 호스트 5433 포트를 사용한다.", "상"],
        ["F_R_243", "권한 구분", "회원 권한은 일반(1)과 최고관리자(10)를 구분해야 한다. 허용되지 않은 권한 값은 수정할 수 없어야 한다.", "상"],
        ["F_R_244", "모델 지연 로드", "서버 기동 시 관리자 페이지가 바로 동작하도록, 분석 모델은 기본으로 첫 분석 요청 때 로드해야 한다.", "상"],
    ]

    grid_table(
        doc,
        ["ID", "요구사항명칭", "설명", "우선순위"],
        func_rows,
        [Cm(2.3), Cm(3.4), Cm(9.2), Cm(1.7)],
    )

    doc.add_page_break()

    # ---- IV. 비기능 ----
    heading(doc, "Ⅳ.  비기능 요구사항")
    add_p(
        doc,
        "적용시점은 알파(핵심 기능 검증), 베타(안정화·품질)로 구분한다.",
        size=10.5,
        align="justify",
        space_after=10,
    )

    nfr_rows = [
        ["NF_R_201", "동시접속", "관리자·학습 API를 포함하여 동시 접속 5명 이상을 처리할 수 있어야 한다.", "베타"],
        ["NF_R_202", "분석 응답속도", "모델이 메모리에 적재된 상태에서 화면 ROI 분류 응답은 네트워크 지연을 제외하고 8초 이내여야 한다. 판독문 생성은 30초 이내를 목표로 한다.", "알파"],
        ["NF_R_203", "화면 전환시간", "관리자 페이지 화면 전환 및 부분(partial) 탭 로딩은 2초 이내여야 한다.", "알파"],
        ["NF_R_204", "업로드 용량", "분석·학습 이미지 1건의 최대 크기는 5MB로 제한해야 한다.", "알파"],
        ["NF_R_205", "가용성", "모델 로드 실패 시에도 로그인·회원·학습 관리 기능은 제공되어야 한다. 분석 API는 503으로 상태를 알려야 한다.", "알파"],
        ["NF_R_206", "신뢰성", "오류는 날짜별 로그 파일에 남기고, 7일 경과 로그는 삭제하여 디스크를 과도하게 사용하지 않아야 한다.", "베타"],
        ["NF_R_207", "지속성", "Docker Compose 환경에서 최소 1개월간 서비스를 재기동 없이 유지할 수 있어야 한다. DB 데이터는 볼륨에 보존되어야 한다.", "베타"],
        ["NF_R_208", "분류 성능", "부위별 의료영상 분류 모델은 학습에 사용한 검증 분포에서 참고용 보조 판독이 가능한 수준이어야 한다. 최종 진단 대체를 요구하지 않는다.", "베타"],
        ["NF_R_209", "사용성", "주요 화면(로그인, 분석, 학습 등록, 회원 관리)은 별도 매뉴얼 없이 사용 가능해야 하며, 사용자 만족도 80점 이상(설문)을 목표로 한다.", "베타"],
        ["NF_R_210", "이식성", "macOS 및 Linux에서 Docker로 실행 가능해야 한다. 추론 디바이스는 CUDA, Apple MPS, CPU 순으로 선택해야 한다. 웹은 Chrome 최신 2개 버전을 지원한다.", "알파"],
        ["NF_R_211", "보안", "학습자 로그인은 세션 쿠키로 유지하고, 비밀번호는 bcrypt 해시를 사용해야 한다. 운영 환경에서는 SECRET_KEY를 기본값에서 변경해야 한다.", "알파"],
        ["NF_R_212", "의료 윤리", "본 시스템은 교육·보조 참고용이다. 모든 분석 결과에 진단 대체 금지 면책 문구를 표시해야 한다.", "알파"],
        ["NF_R_213", "상호운용성", "학습·분석 API는 JSON으로 제공되어 다른 클라이언트(웹 화면)에서 호출할 수 있어야 한다. CORS는 개발 단계에서 허용한다.", "알파"],
        ["NF_R_214", "데이터 무결성", "회원 아이디는 유일해야 한다. 부위는 1~4, 모달리티는 1~3, 채점 결과는 C/H/W, 삭제 여부는 Y/N만 허용해야 한다.", "알파"],
        ["NF_R_215", "실시간성", "SSE keepalive는 30초 이내 간격으로 연결을 유지해야 하며, 클라이언트 종료 시 구독을 해제해야 한다.", "베타"],
        ["NF_R_216", "설정 분리", "DB URL, NAS 계정, Hugging Face 토큰, 시크릿 키는 코드에 하드코딩하지 않고 환경 변수(.env)로 주입해야 한다.", "알파"],
    ]

    grid_table(
        doc,
        ["ID", "요구사항명칭", "설명", "적용시점"],
        nfr_rows,
        [Cm(2.4), Cm(3.2), Cm(9.3), Cm(1.7)],
    )

    add_p(doc, "", size=8, space_after=14)
    add_p(doc, "—  끝  —", size=11, bold=True, align="center", color=NAVY_RGB, space_after=0)

    doc.save(DOCX_PATH)
    print(f"wrote {DOCX_PATH}")
    print(f"diagram {DIAGRAM_PATH}")


if __name__ == "__main__":
    build_diagram()
    build_doc()
